"""Builds a fixture that matches A-CMP's real export: the full COMPRESSOR_FIELDS
column set plus the four derived columns, two sheets, A-CMP v1 format."""
import openpyxl

FIELDS = (
    "id machineTag makeModel serialNo compressorType yearOfManufacture ratedCapacity ratedCapacityUnit "
    "ratedPressure processPressure ratedKw ratedHp ratedRpm ratedCurrent motorEfficiency "
    "annualOperatingDays powerCost starterType designedSec designedAirGen "
    "v1 v2 v3 i1 i2 i3 pf measuredKw kva kvar loadFactor "
    "genLoadVoltage genLoadAmp genLoadPf genLoadKw genUnloadVoltage genUnloadAmp genUnloadPf genUnloadKw "
    "fadActive fadAreaType fadAreaL fadAreaB fadAreaDia fadAreaPeri fadAreaRadius fadAreaDirect "
    "fadNumPoints fadVelocities fadRunningPressure fadMeasuredPower fadSuctionArea fadAvgVelocity "
    "fadAirDeliveryM3Sec fadAirDeliveryM3Hr fadAirDeliveryCfm fadActualSec fadActualAirGen fadDescription "
    "luType luData loadPressure unloadPressure "
    "pumpActive pumpP1 pumpP2 pumpTimeSec pumpAirTempC pumpTempFactor pumpLapData "
    "pumpTankVolume pumpTankVolumeUnit pumpTankCalcMethod pumpTankDia pumpTankLength pumpTankPeri "
    "pumpActualFadM3Min pumpActualFadCfm pumpRunningPressure pumpMeasuredPower pumpDescription "
    "fad operatingPressure inletTemp outletTemp specificPower operatingHours receiverTankPressure noLoadCurrent "
    "photoPath description recordedBy "
    "obsCompSituation obsCompDischarge obsOilSap obsOilRadiatorIn obsOilRadiatorOut "
    "obsAirRadiatorIn obsAirRadiatorOut obsCompFinalDischarge obsCompMotor obsThermalImageNo "
    "createdAt updatedAt createdById"
).split()
DERIVED = ["pumpTankVolumeM3", "luLoadHours", "luUnloadHours", "luTotalHours"]
ALL = FIELDS + DERIVED


def row(**kw):
    d = dict.fromkeys(ALL, "")
    d.update(kw)
    return [d[f] for f in ALL]


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Company Profile"
for r in [
    ["Field", "Value"],
    ["Company Name", "Shree Mahadev Silk Mills Pvt. Ltd."],
    ["Area / Zone", "Kadodara"],
    ["District", "Surat"],
    ["State", "Gujarat"],
    ["Pincode", "394327"],
    ["Overall Consumption (kWh/Month)", 350000],
    ["Exported By", "R. Patel"],
    ["Export Date", "22/09/2026, 11:04:00"],
    ["Format", "A-CMP v1"],
]:
    ws.append(r)

es = wb.create_sheet("Compressor Entries")
es.append(ALL)

# --- AC-01: anemometer traverse, running well above its design SEC
# A choked inlet filter: 4.05 m/s against the ~6.3 m/s this machine
# needs to make its rated 310 CFM, so it is 17 % above its design SEC.
area, vel, kw = 0.0231, 4.05, 41.2
m3s = area * vel
cfm = m3s * 2118.88
es.append(row(
    id="c1", machineTag="AC-01", makeModel="Elgi EG55", serialNo="EG55-2291",
    compressorType="Screw", yearOfManufacture="2018",
    ratedCapacity=310, ratedCapacityUnit="CFM", ratedPressure=7.5, processPressure=6.2,
    ratedKw=55, ratedHp=75, motorEfficiency=94.5, annualOperatingDays=330, starterType="DOL",
    designedSec=round(55 / 310, 4), designedAirGen=round(310 / 55, 3),
    genLoadKw=kw, genUnloadKw=13.8,
    fadActive=True, fadAreaType="Rectangle", fadRunningPressure=6.9, fadMeasuredPower=kw,
    fadSuctionArea=area, fadAvgVelocity=vel,
    fadAirDeliveryM3Sec=round(m3s, 4), fadAirDeliveryCfm=round(cfm, 2),
    fadActualSec=round(kw / cfm, 4), fadActualAirGen=round(cfm / kw, 3),
    fadDescription="Inlet filter differential 0.45 bar, well above the 0.15 bar change-out point.",
    luType="SD", loadPressure=6.2, unloadPressure=7.2,
    description="Aftercooler fins heavily fouled with oil mist.", recordedBy="R. Patel",
    obsCompSituation=38.4, obsCompDischarge=88.6, obsOilSap=74.2,
    obsAirRadiatorIn=91.3, obsAirRadiatorOut=52.8, obsCompMotor=68.9, obsThermalImageNo="IR-114",
    createdAt="2026-09-20T09:12:00Z", updatedAt="2026-09-20T09:40:00Z",
    luLoadHours=4120, luUnloadHours=3180, luTotalHours=7300,
))

# --- AC-02: pump-up test with the temperature correction, idling most of its life
# Healthy on air and power, but loaded only 38 % of its running hours.
p1, p2, tsec, vol_l, temp, kw2 = 4.0, 7.0, 64.0, 2000.0, 41.0, 28.5
tf = 273 / (273 + temp)
vol_m3 = vol_l / 1000
m3min = ((p2 - p1) * vol_m3 * 60 / (1.01325 * tsec)) * tf
cfm2 = m3min * 35.3147
es.append(row(
    id="c2", machineTag="AC-02", makeModel="Atlas Copco GA30", serialNo="GA30-7741",
    compressorType="Screw", yearOfManufacture="2021",
    ratedCapacity=5.1, ratedCapacityUnit="m3/min", ratedPressure=8.0, processPressure=6.0,
    ratedKw=30, motorEfficiency=95.2, annualOperatingDays=330, starterType="VFD",
    genLoadKw=kw2, genUnloadKw=6.1,
    pumpActive=True, pumpP1=p1, pumpP2=p2, pumpTimeSec=tsec, pumpAirTempC=temp,
    pumpTempFactor=round(tf, 4), pumpTankVolume=vol_l, pumpTankVolumeUnit="Liters",
    pumpTankCalcMethod="Direct",
    pumpActualFadM3Min=round(m3min, 4), pumpActualFadCfm=round(cfm2, 2),
    pumpRunningPressure=6.4, pumpMeasuredPower=kw2,
    pumpDescription="Receiver drain valve passing continuously.",
    recordedBy="R. Patel", obsCompDischarge=79.1, obsCompMotor=61.4, obsThermalImageNo="IR-115",
    createdAt="2026-09-20T10:02:00Z", updatedAt="2026-09-20T10:31:00Z",
    pumpTankVolumeM3=vol_m3, luLoadHours=2100, luUnloadHours=3400, luTotalHours=5500,
))

wb.save("t/ACMP_ShreeMahadev.xlsx")

# a second file from another engineer: AC-02 corrected, AC-03 new — tests the merge
wb2 = openpyxl.load_workbook("t/ACMP_ShreeMahadev.xlsx")
e2 = wb2["Compressor Entries"]
e2.cell(row=3, column=ALL.index("pumpMeasuredPower") + 1).value = 25.9
e2.cell(row=3, column=ALL.index("genLoadKw") + 1).value = 25.9
e2.append(row(
    id="c3", machineTag="AC-03", makeModel="Kaeser BSD75", compressorType="Screw",
    yearOfManufacture="2015", ratedCapacity=420, ratedCapacityUnit="CFM",
    ratedPressure=7.0, ratedKw=75, motorEfficiency=93.0, starterType="DOL",
    annualOperatingDays=330, genLoadKw=71.5, genUnloadKw=21.0,
    fadActive=True, fadSuctionArea=0.0290, fadAvgVelocity=7.1, fadMeasuredPower=71.5,
    fadRunningPressure=6.8, fadAirDeliveryCfm=round(0.0290 * 7.1 * 2118.88, 2),
    recordedBy="H. Patel", createdAt="2026-09-21T08:00:00Z", updatedAt="2026-09-21T08:30:00Z",
    luLoadHours=5600, luUnloadHours=900, luTotalHours=6500,
))
wb2.save("t/ACMP_SecondEngineer.xlsx")

print("wrote t/ACMP_ShreeMahadev.xlsx and t/ACMP_SecondEngineer.xlsx")
print("AC-01 expect  rated 310.0 CFM · actual %.2f · SEC %.4f vs design %.4f · %+.1f %%"
      % (cfm, kw / cfm, 55 / 310, ((kw / cfm) - (55 / 310)) / (55 / 310) * 100))
print("AC-02 expect  rated %.1f CFM · actual %.2f (temp factor %.4f) · loaded %.1f %%"
      % (5.1 * 35.3147, cfm2, tf, 2100 / 5500 * 100))
